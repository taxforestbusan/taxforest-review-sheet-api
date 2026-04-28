"""
종합소득세 검토표 자동생성기 - FastAPI 백엔드
Railway 배포용

엔드포인트:
  POST /api/generate-review-sheet
    multipart/form-data:
      - prev_sinjako: PDF (전년 귀속 신고서, 필수)
      - prev_annae:   PDF (전년 귀속 안내문, 선택)
      - curr_annae:   PDF (당해 귀속 안내문, 선택)
      - curr_susip_manual: int (당해 수입금액 수동입력, 선택)
    응답: 검토표 XLSX 파일 (binary)

  POST /api/parse-only
    위와 동일한 입력, 응답: JSON (파싱 결과만 반환 - 미리보기용)
"""
import io
import re
import tempfile
import os
from datetime import datetime
from typing import Optional

import pdfplumber
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, JSONResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter


app = FastAPI(title="종합소득세 검토표 생성기 API")

# CORS - Netlify 도메인 허용
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # 운영 시 ["https://your-netlify-app.netlify.app"]
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ====================================================================
# 1) 파서: 종합소득세 신고서 (전자신고제출분)
# ====================================================================

def _to_int(s):
    if s is None:
        return None
    s = str(s).replace(',', '').replace(' ', '').strip()
    try:
        return int(s)
    except ValueError:
        return None


# 신고서 1페이지 4. 세액의 계산 라벨 매핑
# (키, 종소세번호, 농특세번호, 라벨 정규식)
SECT4_ITEMS = [
    ('jonghap_sodeuk',     19, None,  r'종\s*합\s*소\s*득\s*금\s*액'),
    ('sodeuk_gongje',      20, None,  r'^소\s*득\s*공\s*제'),
    ('gwasea_pyojun',      21, 41,    r'과\s*세\s*표\s*준'),
    ('serul',              22, 42,    r'^세\s*율'),
    ('sanchul_seak',       23, 43,    r'^산\s*출\s*세\s*액'),
    ('seak_gammyeon',      24, None,  r'^세\s*액\s*감\s*면'),
    ('seak_gongje',        25, None,  r'^세\s*액\s*공\s*제'),
    ('jonghap_gwasea',     26, 44,    r'^종\s*합\s*과\s*세'),
    ('hapge_28',           28, 46,    r'^합\s*계\(\s*2\s*6'),
    ('gasanseo',           29, 47,    r'^가\s*산\s*세'),
    ('hapge_31',           31, 49,    r'^합\s*계\(\s*2\s*8\s*\+\s*2\s*9'),
    ('gi_napbu',           32, 50,    r'^기\s*납\s*부\s*세\s*액'),
    ('napbu_total_33',     33, 51,    r'^납\s*부.*할\s*총\s*세\s*액'),
    ('shingo_napbu_37',    37, 53,    r'^신고기한이내납부'),
]


def _parse_line_with_seq(line: str, seq_jong: int, seq_nong):
    """라인에서 종소세번호와 농특세번호를 기준으로 종소/농특 값 분리 추출"""
    pat_seq_jong = rf'(?<![0-9])\s{seq_jong}\s'
    m_jong = re.search(pat_seq_jong, line)
    if not m_jong:
        return None, None
    after_jong = line[m_jong.end():]
    val_jong, val_nong = None, None
    if seq_nong is not None:
        pat_seq_nong = rf'(?<![0-9])\s*{seq_nong}\s*$|(?<![0-9])\s*{seq_nong}\s'
        m_nong = re.search(pat_seq_nong, after_jong)
        if m_nong:
            jong_zone = after_jong[:m_nong.start()]
            nong_zone = after_jong[m_nong.end():]
        else:
            jong_zone, nong_zone = after_jong, ''
    else:
        jong_zone, nong_zone = after_jong, ''
    nums_jong = re.findall(r'[\d,]+', jong_zone.strip())
    nums_nong = re.findall(r'[\d,]+', nong_zone.strip())
    if nums_jong:
        val_jong = _to_int(nums_jong[0])
    if nums_nong:
        val_nong = _to_int(nums_nong[0])
    return val_jong, val_nong


def parse_sinjuk_sinjako(pdf_bytes: bytes) -> dict:
    """종합소득세 신고서 PDF 파싱"""
    result = {'guisok_year': None}
    
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        page1 = pdf.pages[0].extract_text() or ''
        page2 = pdf.pages[1].extract_text() if len(pdf.pages) > 1 else ''
    
    # 귀속연도
    m = re.search(r'\(\s*([\d\s]{4,8})년\s*귀속', page1)
    if m:
        y = m.group(1).replace(' ', '')
        if len(y) == 4:
            result['guisok_year'] = int(y)
    
    # 4. 세액의 계산
    lines = page1.split('\n')
    sect4_start = 0
    for i, line in enumerate(lines):
        if re.search(r'4\s+세액의\s+계산', line):
            sect4_start = i
            break
    sect4_lines = lines[sect4_start:sect4_start + 30]
    
    for key, seq_j, seq_n, label_pat in SECT4_ITEMS:
        for line in sect4_lines:
            line_stripped = line.strip()
            if re.search(label_pat, line_stripped):
                v_j, v_n = _parse_line_with_seq(line_stripped, seq_j, seq_n)
                result[key] = v_j
                if seq_n is not None:
                    result[key + '_nong'] = v_n
                break
    
    # 사업소득명세서 (2페이지)
    if page2:
        m = re.search(r'⑨\s*총\s*수\s*입\s*금\s*액\s+([\d,\s]+)', page2)
        if m:
            nums = re.findall(r'[\d,]+', m.group(1))
            if nums:
                result['sa_chong_susip'] = _to_int(nums[-1])
        m = re.search(r'⑩\s*필\s*요\s*경\s*비\s+([\d,\s]+)', page2)
        if m:
            nums = re.findall(r'[\d,]+', m.group(1))
            if nums:
                result['sa_pilyo_gyongbi'] = _to_int(nums[-1])
        m = re.search(r'⑪\s*소\s*득\s*금\s*액', page2)
        if m:
            after = page2[m.end():].split('\n')[0]
            nums = re.findall(r'[\d,]+', after)
            if nums:
                result['sa_sodeuk_geum'] = _to_int(nums[-1])
        m = re.search(r'④\s*상\s*호\s+(\S+)', page2)
        if m:
            result['sa_sangho'] = m.group(1)
        m = re.search(r'⑤\s*사\s*업\s*자\s*등\s*록\s*번\s*호\s+(\d{3}-\d{2}-\d{5})', page2)
        if m:
            result['sa_saupja_no'] = m.group(1)
        m = re.search(r'⑧\s*주\s*업\s*종\s*코\s*드\s+(\d+)', page2)
        if m:
            result['sa_eopjong_code'] = m.group(1)
    
    return result


# ====================================================================
# 2) 파서: 종합소득세 신고 안내문
# ====================================================================

def parse_annae(pdf_bytes: bytes) -> dict:
    """안내문 PDF 파싱"""
    result = {
        'guisok_year': None, 'sangho': None, 'saupja_no': None,
        'name': None, 'birth': None,
        'annae_yuhyung': None, 'gijang_uimu': None, 'chuge_gyongbi': None,
        'eopjong_code': None, 'sa_susip': None,
        'gijun_gb_normal': None, 'gijun_gb_jaga': None,
        'dansoon_gb_normal': None, 'dansoon_gb_jaga': None,
        'gukmin_yeongeum': None, 'noran_usan': None,
        'gaein_yeongeum': None, 'tweejik_yeongeum': None, 'yeongeum_gyejwa': None,
        'last3_jonghap_sodeuk': None, 'last3_gwasea': None,
        'last3_gyuljeong': None, 'last3_silhyo_serul': None,
        'last3_susip': None, 'last3_pilyo': None, 'last3_sodeukrul': None,
        'card_total_amount': None, 'card_eopmu_mugwan_amount': None,
        'card_chiryo_amount': None, 'card_gajeong_amount': None,
        'card_sinbyun_amount': None, 'card_haeoe_amount': None,
        'individual_analysis': None,
    }
    
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        full = '\n'.join(p.extract_text() or '' for p in pdf.pages)
    
    m = re.search(r'안내\s*정보\s*\(\s*(\d{4})\s*년\s*귀속', full)
    if m:
        result['guisok_year'] = int(m.group(1))
    m = re.search(r'성명\s+(\S+)\s+생년월일\s+([\d.]+)', full)
    if m:
        result['name'] = m.group(1)
        result['birth'] = m.group(2)
    m = re.search(r'안내유형\s+(\S+)', full)
    if m: result['annae_yuhyung'] = m.group(1)
    m = re.search(r'기장의무\s+(\S+)', full)
    if m: result['gijang_uimu'] = m.group(1)
    m = re.search(r'추계시\s*적용경비율\s+(\S+)', full)
    if m: result['chuge_gyongbi'] = m.group(1)
    
    m = re.search(r'총\s*계\s*([\d,]+(?:\s*\n\s*\d+)?)', full)
    if m:
        raw = re.sub(r'[\s,]+', '', m.group(1))
        result['sa_susip'] = _to_int(raw)
    
    m = re.search(r'(\d{3}-\d{2}-\d{5})(\S+)\s+\S+\s+(\d{6})', full)
    if m:
        result['saupja_no'] = m.group(1)
        result['sangho'] = m.group(2)
        result['eopjong_code'] = m.group(3)
    
    m = re.search(r'(\d{6})[\s\S]*?([\d.]+)\s*%\s+([\d.]+)\s*%\s+([\d.]+)\s*%\s+([\d.]+)\s*%', full)
    if m:
        result['gijun_gb_normal'] = float(m.group(2))
        result['gijun_gb_jaga'] = float(m.group(3))
        result['dansoon_gb_normal'] = float(m.group(4))
        result['dansoon_gb_jaga'] = float(m.group(5))
    
    m = re.search(r'국민연금보험료\s+([\d,]+)\s*원', full)
    if m: result['gukmin_yeongeum'] = _to_int(m.group(1))
    m = re.search(r'개인연금저축\s+([\d,]+)\s*원', full)
    if m: result['gaein_yeongeum'] = _to_int(m.group(1))
    m = re.search(r'소기업소상공인공제부금[\s\S]{0,200}?([\d,]+)\s*원', full)
    if m: result['noran_usan'] = _to_int(m.group(1))
    m = re.search(r'퇴직연금세액공제\s+([\d,]+)\s*원', full)
    if m: result['tweejik_yeongeum'] = _to_int(m.group(1))
    m = re.search(r'연금계좌세액공제\s+([\d,]+)\s*원', full)
    if m: result['yeongeum_gyejwa'] = _to_int(m.group(1))
    
    m = re.search(r'종합소득금액\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)', full)
    if m: result['last3_jonghap_sodeuk'] = _to_int(m.group(3)) * 1000
    m = re.search(r'결정세액\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)', full)
    if m: result['last3_gyuljeong'] = _to_int(m.group(3)) * 1000
    m = re.search(r'과세표준\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)', full)
    if m: result['last3_gwasea'] = _to_int(m.group(3)) * 1000
    m = re.search(r'실효세율\s+([\d.]+)\s*%\s+([\d.]+)\s*%\s+([\d.]+)\s*%', full)
    if m: result['last3_silhyo_serul'] = float(m.group(3))
    
    m = re.search(r'수입금액\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)', full)
    if m: result['last3_susip'] = _to_int(m.group(3)) * 1000
    m = re.search(r'필요경비\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)', full)
    if m: result['last3_pilyo'] = _to_int(m.group(3)) * 1000
    m = re.search(r'소득률\(당해업체\)\s+([\d.]+)\s*%\s+([\d.]+)\s*%\s+([\d.]+)\s*%', full)
    if m: result['last3_sodeukrul'] = float(m.group(3))
    
    m = re.search(r'금액\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)', full)
    if m:
        result['card_total_amount'] = _to_int(m.group(1))
        result['card_sinbyun_amount'] = _to_int(m.group(2))
        result['card_gajeong_amount'] = _to_int(m.group(3))
        result['card_eopmu_mugwan_amount'] = _to_int(m.group(4))
        result['card_chiryo_amount'] = _to_int(m.group(5))
        result['card_haeoe_amount'] = _to_int(m.group(6))
    
    m = re.search(r'소득세\s*개별분석자료\s*[:：]\s*([^\n]+)', full)
    if m: result['individual_analysis'] = m.group(1).strip()
    
    return result


# ====================================================================
# 3) 검토표 엑셀 생성
# ====================================================================

# 셀 스타일
THIN = Side(style='thin', color='000000')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill('solid', start_color='4472C4')
SUBHEADER_FILL = PatternFill('solid', start_color='D9E1F2')
DIFF_FILL = PatternFill('solid', start_color='FFF2CC')


def build_review_workbook(prev_s: dict, prev_a: dict, curr_a: dict,
                           curr_susip_manual: Optional[int] = None) -> Workbook:
    """검토표 엑셀 워크북 생성 (템플릿 없이 from scratch)"""
    wb = Workbook()
    ws = wb.active
    ws.title = '검토표'
    
    prev_year = prev_s.get('guisok_year')
    if curr_a and curr_a.get('guisok_year'):
        curr_year = curr_a['guisok_year']
    elif prev_year:
        curr_year = prev_year + 1
    else:
        curr_year = None
    
    sangho = (prev_s.get('sa_sangho') or prev_a.get('sangho') 
              or (curr_a.get('sangho') if curr_a else None) or '')
    name = prev_a.get('name') or (curr_a.get('name') if curr_a else None) or ''
    
    # === 1행: 제목 ===
    ws.merge_cells('A1:E1')
    ws['A1'] = '종합소득세 검토표'
    ws['A1'].font = Font(name='맑은 고딕', size=16, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill('solid', start_color='1F4E79')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 32
    
    # === 3행: 상호 ===
    ws['A3'] = '상호:'
    ws['A3'].font = Font(name='맑은 고딕', size=11, bold=True)
    ws['B3'] = f'{sangho} ({name})' if name and sangho else sangho or name
    ws.merge_cells('B3:C3')
    ws['B3'].font = Font(name='맑은 고딕', size=11)
    
    # === 5행: 헤더 ===
    headers = {
        'A5': '구분',
        'B5': f'{prev_year}년 귀속' if prev_year else '전년 귀속',
        'C5': f'{curr_year}년 귀속' if curr_year else '당해 귀속',
        'D5': '비고',
        'E5': '메모',
    }
    for cell, val in headers.items():
        ws[cell] = val
        ws[cell].font = Font(name='맑은 고딕', size=10, bold=True, color='FFFFFF')
        ws[cell].fill = HEADER_FILL
        ws[cell].alignment = Alignment(horizontal='center', vertical='center')
        ws[cell].border = BOX
    
    # === 데이터 영역 ===
    # 신고유형
    yuhyung_b = ''
    if prev_a.get('annae_yuhyung'):
        yuhyung_b = prev_a['annae_yuhyung']
        if prev_a.get('gijang_uimu'):
            yuhyung_b += f" / {prev_a['gijang_uimu']}"
    
    yuhyung_c = ''
    if curr_a and curr_a.get('annae_yuhyung'):
        yuhyung_c = curr_a['annae_yuhyung']
        if curr_a.get('gijang_uimu'):
            yuhyung_c += f" / {curr_a['gijang_uimu']}"
    
    ws['A6'] = '신고유형'
    ws['B6'] = yuhyung_b
    ws['C6'] = yuhyung_c
    
    # 총수입금액 (A7~A12 병합 - 사업장 추가용 여유)
    ws.merge_cells('A7:A12')
    ws['A7'] = '총수입금액'
    ws['B7'] = prev_s.get('sa_chong_susip') or prev_a.get('sa_susip')
    if curr_susip_manual is not None:
        ws['C7'] = curr_susip_manual
    elif curr_a:
        ws['C7'] = curr_a.get('sa_susip')
    
    # 합계 행 (12)
    ws['B12'] = '=SUM(B7:B11)'
    ws['C12'] = '=SUM(C7:C11)'
    ws['D12'] = '합계'
    
    # 기준경비율
    ws['A13'] = '기준경비율'
    if prev_a.get('gijun_gb_normal') is not None:
        ws['B13'] = prev_a['gijun_gb_normal'] / 100
        ws['B13'].number_format = '0.0%'
    if curr_a and curr_a.get('gijun_gb_normal') is not None:
        ws['C13'] = curr_a['gijun_gb_normal'] / 100
        ws['C13'].number_format = '0.0%'
    
    # 신고소득률 (자동계산)
    ws['A14'] = '신고소득률'
    ws['B14'] = '=IF(B12=0,0,B15/B12)'
    ws['B14'].number_format = '0.00%'
    ws['C14'] = '=IF(C12=0,0,C15/C12)'
    ws['C14'].number_format = '0.00%'
    
    # 신고소득금액
    ws['A15'] = '신고소득금액'
    ws['B15'] = prev_s.get('sa_sodeuk_geum') or prev_s.get('jonghap_sodeuk')
    
    # 결정소득금액 (종합소득금액)
    ws['A16'] = '결정소득금액'
    ws['B16'] = prev_s.get('jonghap_sodeuk')
    
    # 소득공제
    ws['A17'] = '소득공제'
    ws['B17'] = prev_s.get('sodeuk_gongje')
    
    # 과세표준
    ws['A18'] = '과세표준'
    ws['B18'] = prev_s.get('gwasea_pyojun')
    
    # 세율
    ws['A19'] = '세율'
    if prev_s.get('serul') is not None:
        ws['B19'] = prev_s['serul'] / 100
        ws['B19'].number_format = '0%'
    
    # 산출세액
    ws['A20'] = '산출세액'
    ws['B20'] = prev_s.get('sanchul_seak')
    
    # 세액감면
    ws['A21'] = '세액감면'
    ws['B21'] = prev_s.get('seak_gammyeon') or 0
    
    # 세액공제
    ws['A22'] = '세액공제'
    ws['B22'] = prev_s.get('seak_gongje') or 0
    
    # 결정세액 (수식)
    ws['A23'] = '결정세액'
    ws['B23'] = '=B20-B21-B22'
    
    # 가산세
    ws['A24'] = '가산세'
    ws['B24'] = prev_s.get('gasanseo') or 0
    
    # 추가납부세액
    ws['A25'] = '추가납부세액'
    ws['B25'] = 0
    
    # 기납부세액
    ws['A26'] = '기납부세액'
    ws['B26'] = prev_s.get('gi_napbu') or 0
    
    # 납부세액 (수식)
    ws['A27'] = '납부세액'
    ws['B27'] = '=B23+B24+B25-B26'
    
    # 농어촌특별세 (1페이지 농특세 합계 또는 별도 라인)
    ws['A28'] = '농어촌특별세'
    nong = prev_s.get('hapge_31_nong') or prev_s.get('jonghap_gwasea_nong') or 0
    ws['B28'] = nong
    
    # 익금/손금가산
    ws['A29'] = '익금가산금액'
    ws['B29'] = 0
    ws['A30'] = '손금가산금액'
    ws['B30'] = 0
    
    # === 메모 영역 (E열) ===
    memo_lines = []
    memo_lines.append(('∨ 타소득', True))
    memo_lines.append(('· 주택임대소득 - ', False))
    memo_lines.append(('', False))
    memo_lines.append(('', False))
    memo_lines.append(('∨ 인적공제', True))
    memo_lines.append(('· 본인', False))
    memo_lines.append(('· ', False))
    memo_lines.append(('· ', False))
    memo_lines.append(('', False))
    memo_lines.append(('∨ 기부금 - ', True))
    memo_lines.append(('', False))
    memo_lines.append(('∨ 공제/감면', True))
    memo_lines.append(('· 중소기업특별세액감면', False))
    memo_lines.append(('· ', False))
    memo_lines.append(('', False))
    memo_lines.append(('', False))
    memo_lines.append(('∨ 지방세 - ', True))
    
    for i, (text, is_header) in enumerate(memo_lines):
        cell = ws.cell(row=6 + i, column=5, value=text)
        cell.font = Font(name='맑은 고딕', size=9, bold=is_header)
    
    # === 자동 분석 메모 (D열 비고) ===
    notes = []
    if prev_a.get('individual_analysis'):
        notes.append((f"개별분석: {prev_a['individual_analysis']}", 19))
    if prev_a.get('card_eopmu_mugwan_amount'):
        total = prev_a.get('card_total_amount', 1) or 1
        pct = prev_a['card_eopmu_mugwan_amount'] / total * 100
        notes.append((
            f"업무무관 카드: {prev_a['card_eopmu_mugwan_amount']:,}원 ({pct:.1f}%)",
            20
        ))
    if prev_a.get('card_gajeong_amount'):
        notes.append((f"가정용품 카드: {prev_a['card_gajeong_amount']:,}원", 21))
    if prev_a.get('card_chiryo_amount'):
        notes.append((f"개인치료 카드: {prev_a['card_chiryo_amount']:,}원", 22))
    if prev_a.get('card_sinbyun_amount'):
        notes.append((f"신변잡화 카드: {prev_a['card_sinbyun_amount']:,}원", 23))
    if prev_a.get('noran_usan'):
        notes.append((f"노란우산: {prev_a['noran_usan']:,}원", 24))
    
    for text, row in notes:
        cell = ws.cell(row=row, column=4, value=text)
        cell.font = Font(name='맑은 고딕', size=9, color='C00000')
    
    # === 셀 스타일 일괄 적용 ===
    for row in range(6, 31):
        for col in range(1, 6):
            cell = ws.cell(row=row, column=col)
            cell.border = BOX
            if col == 1:  # A열 (구분)
                cell.font = Font(name='맑은 고딕', size=10, bold=True)
                cell.fill = SUBHEADER_FILL
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col in (2, 3):  # B,C열 (값)
                if cell.font.name != '맑은 고딕':
                    cell.font = Font(name='맑은 고딕', size=10)
                cell.alignment = Alignment(horizontal='right', vertical='center')
                # 숫자 포맷 (퍼센트 셀이 아닌 경우만)
                if cell.number_format == 'General' and isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0;(#,##0);-'
            elif col == 4:  # D열 (비고)
                if cell.font.name != '맑은 고딕':
                    cell.font = Font(name='맑은 고딕', size=9)
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            elif col == 5:  # E열 (메모)
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    
    # 컬럼 너비
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 32
    ws.column_dimensions['E'].width = 30
    
    # 합계 행 강조
    for col in ('B12', 'C12'):
        ws[col].fill = SUBHEADER_FILL
        ws[col].font = Font(name='맑은 고딕', size=10, bold=True)
    ws['D12'].fill = SUBHEADER_FILL
    ws['D12'].font = Font(name='맑은 고딕', size=10, bold=True)
    
    # 결정세액/납부세액 행 강조
    for row in (23, 27):
        for col in range(1, 4):
            ws.cell(row=row, column=col).fill = DIFF_FILL
            ws.cell(row=row, column=col).font = Font(name='맑은 고딕', size=10, bold=True)
    
    return wb


# ====================================================================
# 4) FastAPI 엔드포인트
# ====================================================================

@app.get("/")
def health():
    return {"status": "ok", "service": "종합소득세 검토표 생성기"}


@app.post("/api/parse-only")
async def parse_only(
    prev_sinjako: UploadFile = File(...),
    prev_annae: Optional[UploadFile] = File(None),
    curr_annae: Optional[UploadFile] = File(None),
):
    """파싱 결과를 JSON으로 반환 (미리보기용)"""
    try:
        prev_s_bytes = await prev_sinjako.read()
        prev_s = parse_sinjuk_sinjako(prev_s_bytes)
        
        prev_a = {}
        if prev_annae:
            prev_a = parse_annae(await prev_annae.read())
        
        curr_a = {}
        if curr_annae:
            curr_a = parse_annae(await curr_annae.read())
        
        return JSONResponse({
            "prev_sinjako": prev_s,
            "prev_annae": prev_a,
            "curr_annae": curr_a,
        })
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"파싱 오류: {str(e)}")


@app.post("/api/generate-review-sheet")
async def generate_review_sheet_endpoint(
    prev_sinjako: UploadFile = File(...),
    prev_annae: Optional[UploadFile] = File(None),
    curr_annae: Optional[UploadFile] = File(None),
    curr_susip_manual: Optional[int] = Form(None),
):
    """검토표 엑셀 파일을 생성하여 반환"""
    try:
        prev_s = parse_sinjuk_sinjako(await prev_sinjako.read())
        prev_a = parse_annae(await prev_annae.read()) if prev_annae else {}
        curr_a = parse_annae(await curr_annae.read()) if curr_annae else {}
        
        wb = build_review_workbook(prev_s, prev_a, curr_a, curr_susip_manual)
        
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        
        # 파일명: {상호}_종소세검토표_{귀속}.xlsx
        sangho = (prev_s.get('sa_sangho') or prev_a.get('sangho') 
                  or (curr_a.get('sangho') if curr_a else None) or '검토표')
        prev_year = prev_s.get('guisok_year') or 0
        curr_year = (curr_a.get('guisok_year') if curr_a else None) or (prev_year + 1)
        filename = f"{sangho}_종소세검토표_{curr_year}귀속.xlsx"
        
        # 한글 파일명 인코딩 (RFC 5987)
        from urllib.parse import quote
        filename_encoded = quote(filename)
        
        return StreamingResponse(
            buf,
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={
                "Content-Disposition": f"attachment; filename*=UTF-8''{filename_encoded}"
            }
        )
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=400, detail=f"생성 오류: {str(e)}")


if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", 8000))
    uvicorn.run(app, host="0.0.0.0", port=port)
